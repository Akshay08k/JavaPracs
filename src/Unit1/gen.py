import os
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
fileName = "MCA_A_22_JAVA_UNIT1.docx"
def apply_code_style(paragraph):
    """Apply monospaced styling to a paragraph."""
    run = paragraph.runs[0]
    run.font.name = "Courier New"
    run.font.size = Pt(10)
    shading_elm = parse_xml(r'<w:shd {} w:fill="EDEDED"/>'.format(nsdecls('w')))
    paragraph._element.get_or_add_pPr().append(shading_elm)

def create_word_from_java_files():
    doc = Document()
    doc.add_heading("Java Code Files", level=1)
    
    for file in os.listdir():
        if file.endswith(".java"):
            doc.add_heading(file, level=2)
            with open(file, "r", encoding="utf-8") as f:
                code_content = f.read()
                paragraph = doc.add_paragraph(code_content)
                apply_code_style(paragraph)
                doc.add_paragraph("\n")  # Add spacing
    
    doc.save(fileName)
    print(f"Word file created: {fileName}")

if __name__ == "__main__":
    create_word_from_java_files()
